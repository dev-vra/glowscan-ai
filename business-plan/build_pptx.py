"""Gera pitch-vico.pptx (16:9, identidade Viço) a partir de model.py + data.json + charts/."""
import json

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Pt

import model

TERRA, INK, MUTED, CREME, SURF, SOL = (RGBColor(0xB8, 0x46, 0x2F), RGBColor(0x2A, 0x1B, 0x14), RGBColor(0x6E, 0x5C, 0x51),
                                       RGBColor(0xFF, 0xFB, 0xF8), RGBColor(0xF6, 0xEE, 0xE7), RGBColor(0xF2, 0xB6, 0x6D))
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
data = json.load(open("data.json", encoding="utf-8"))
ys = model.yearly(model.run())
ue = model.unit_economics()
S = data["scenarios"]
RAISE = 1_500_000
def mi(v):
    sign, a = ("-" if v < 0 else ""), abs(v)
    return f"{sign}R$ {a / 1e6:.1f} mi".replace(".", ",") if a >= 1e6 else f"{sign}R$ {a / 1e3:.0f} mil"
n = lambda v: f"{v:,.0f}".replace(",", ".")

prs = Presentation()
prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
W, Hh = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def slide(bg=CREME):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    return s


def text(s, x, y, w, h, t, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    lines = t if isinstance(t, list) else [t]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold, color, "Calibri"
        p.space_after = Pt(size * 0.5)
    return tb


def card(s, x, y, w, h, fill=WHITE):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    sh.adjustments[0] = 0.12
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def title(s, eyebrow, t):
    text(s, 2, 1.4, 30, 1, eyebrow.upper(), 13, TERRA, True)
    text(s, 2, 2.2, 30, 2.4, t, 32, INK, True)


def footer(s, i):
    text(s, 2, 17.7, 10, 0.8, "viço", 12, TERRA, True)
    text(s, 28.5, 17.7, 3.5, 0.8, str(i), 11, MUTED, align=PP_ALIGN.RIGHT)


def stat_cards(s, items, y=6.2, h=5):
    w = (29.9 - 0.8 * (len(items) - 1)) / len(items)
    for i, (big, small) in enumerate(items):
        x = 2 + i * (w + 0.8)
        card(s, x, y, w, h)
        text(s, x + 0.7, y + 0.7, w - 1.4, 2.2, big, 34, TERRA, True)
        text(s, x + 0.7, y + 2.8, w - 1.4, 2, small, 14, INK)


k = 0
# 1. Capa
s = slide(TERRA)
text(s, 2.2, 4.5, 28, 4, "viço", 120, WHITE, True)
text(s, 2.4, 10.4, 28, 2, "Descubra se o seu skincare está funcionando.", 30, WHITE, True)
text(s, 2.4, 12.6, 28, 1.5, f"Pré-seed · {mi(RAISE)} · Confidencial", 16, RGBColor(0xFB, 0xED, 0xD6))

# 2. Problema
k += 2
s = slide()
title(s, "O problema", "Quem compra skincare não sabe se está funcionando")
stat_cards(s, [("3 apps", "um para a pele, outro para o rótulo, outro para anotar — ninguém liga os três"),
               ("216", "avaliações de 9 apps concorrentes: câmera falha, paywall sem valor, resultado inconsistente"),
               ("Só a marca", "ferramentas de grandes marcas recomendam apenas o próprio portfólio")])
footer(s, k)

# 3. Solução
k += 1
s = slide()
title(s, "A solução", "Selfie, rótulo e rotina — num ciclo que mede o resultado")
steps = [("1", "Lê a pele", "Skin Score 0–100 e 7 métricas por zona, com câmera que só dispara em boa luz."),
         ("2", "Lê os produtos", "Foto do rótulo; a IA entende cada ingrediente, inclusive de farmácia brasileira."),
         ("3", "Monta a rotina", "Manhã e noite na ordem certa, conflitos corrigidos em um toque."),
         ("4", "Mede a evolução", "Gráfico, antes/depois e card para compartilhar — sem mostrar o rosto.")]
for i, (num, t, d) in enumerate(steps):
    x = 2 + i * 7.6
    card(s, x, 6.2, 7, 9)
    text(s, x + 0.6, 6.8, 2, 1.6, num, 36, TERRA, True)
    text(s, x + 0.6, 8.8, 6, 1.2, t, 20, INK, True)
    text(s, x + 0.6, 10.3, 6, 4.5, d, 14, MUTED)
footer(s, k)

# 4. Produto no ar
k += 1
s = slide()
title(s, "Tração de produto", "Construído e no ar — não é um slide")
text(s, 2, 6.2, 15, 10, ["• Beta fechado em vico.bitrilha.com.br", "• App Android empacotado; iOS na sequência",
                         "• IA de visão para pele e rótulos", "• Motor de rotina com 12 regras de conflito, testado",
                         "• Pagamentos, LGPD e painel de testers prontos"], 18, INK)
card(s, 18.5, 6.2, 13.4, 9, TERRA)
text(s, 19.3, 7, 12, 2, "Próximo marco", 14, RGBColor(0xFB, 0xED, 0xD6), True)
text(s, 19.3, 8.4, 12, 6, "Conversão teste→pago ≥ 30% medida com os primeiros testers e lançamento público nas lojas.", 22, WHITE, True)
footer(s, k)

# 5. Mercado
k += 1
s = slide()
title(s, "Mercado", "3º maior mercado de beleza do mundo, e a categoria certa crescendo")
stat_cards(s, [("R$ 200 bi", "beleza e cuidados pessoais no Brasil em 2025 (ABIHPEC)"),
               ("R$ 12,3 bi", "dermocosméticos em 12 meses, +8,3% (IQVIA)"),
               ("US$ 1,8 bi", "apps de análise de pele no mundo, 17% ao ano (FMI)"),
               ("23%", "dos brasileiros já usam IA para escolher beleza (out/2026)")])
text(s, 2, 12.4, 30, 3, "Prova de mecânica: o Cal AI aplicou foto → IA → acompanhamento a calorias e chegou a ~US$ 30 mi de ARR sem investimento externo, até ser comprado pela MyFitnessPal.", 15, MUTED)
footer(s, k)

# 6. Concorrência
k += 1
s = slide()
title(s, "Concorrência", "O único que fecha o ciclo — e é neutro")
rows = [["", "Pele", "Rótulo", "Rotina", "Evolução", "Neutro"],
        ["Viço", "●", "●", "●", "●", "●"], ["Skan", "●", "", "◐", "◐", ""], ["Yuka", "", "●", "", "", "●"],
        ["FeelinMySkin", "", "◐", "◐", "●", "●"], ["Marcas (Natura, Mary Kay)", "●", "", "◐", "", ""]]
tbl = s.shapes.add_table(len(rows), 6, Cm(2), Cm(6), Cm(29.9), Cm(9)).table
for r, row in enumerate(rows):
    for c, v in enumerate(row):
        cell = tbl.cell(r, c)
        cell.text = v
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
        run = p.runs[0] if p.runs else p.add_run()
        run.font.size = Pt(16)
        run.font.bold = r == 0 or (r == 1)
        run.font.color.rgb = WHITE if r == 0 else (TERRA if r == 1 else INK)
        cell.fill.solid()
        cell.fill.fore_color.rgb = TERRA if r == 0 else (RGBColor(0xF9, 0xE6, 0xDE) if r == 1 else WHITE)
tbl.columns[0].width = Cm(9.9)
for c in range(1, 6):
    tbl.columns[c].width = Cm(4)
footer(s, k)

# 7. Modelo
k += 1
s = slide()
title(s, "Modelo de negócio", "Assinatura no B2C, dados e parcerias no B2B")
card(s, 2, 6.2, 14.5, 9.5)
text(s, 2.8, 6.8, 13, 1.2, "B2C — assinatura", 20, TERRA, True)
text(s, 2.8, 8.3, 13, 7, ["R$ 24,90/mês ou R$ 199/ano", "7 dias grátis, cancelamento em um toque",
                          f"LTV de contribuição ~R$ {ue['ltv']:.0f}", f"CAC misto ano 1 ~R$ {ys[0]['cac'] / (1 + model.Y['organic_ratio'][0]):.0f}"], 17)
card(s, 17.4, 6.2, 14.5, 9.5)
text(s, 18.2, 6.8, 13, 1.2, "B2B — a partir do ano 2", 20, TERRA, True)
text(s, 18.2, 8.3, 13, 7, ["Afiliados de recompra", "Estudos de eficácia patrocinados", "Painel de tendências anonimizado",
                           "White-label para redes de farmácia"], 17)
footer(s, k)

# 8. Go-to-market
k += 1
s = slide()
title(s, "Go-to-market", "Conteúdo que se espalha sozinho")
stat_cards(s, [("TikTok/Reels", "card \"minha pele em 30 dias\" e antes/depois — o produto gera o conteúdo"),
               ("Micro-influência", "criadoras de skincare de 5–20 mil seguidores, por permuta e afiliado"),
               ("SEO de ativos", "\"retinol e ácido pode junto?\" — páginas geradas da base de regras")])
footer(s, k)

# 9. Projeções
k += 1
s = slide()
title(s, "Projeções (cenário base)", f"{mi(ys[4]['gross'])} de receita e {n(ys[4]['subs_end'])} assinantes no ano 5")
s.shapes.add_picture("charts/receita.png", Cm(2), Cm(5.6), width=Cm(16.5))
s.shapes.add_picture("charts/caixa.png", Cm(18.8), Cm(5.6), width=Cm(13.2))
text(s, 2, 15.2, 30, 2, "Modelo mensal de 60 meses com fórmulas abertas. Teste→pago 40% (mediana Saúde & Fitness, RevenueCat 2025), churn mensal 12%, renovação anual 35%.", 12, MUTED)
footer(s, k)

# 10. Cenários
k += 1
s = slide()
title(s, "Cenários", "Plano robusto mesmo no conservador")
s.shapes.add_picture("charts/cenarios.png", Cm(2), Cm(5.8), width=Cm(17))
text(s, 20, 6.2, 12, 10, [f"Conservador: {n(S['Conservador']['subs5'])} assinantes, {mi(S['Conservador']['gross5'])} no ano 5",
                          f"Base: {n(S['Base']['subs5'])} assinantes, {mi(S['Base']['gross5'])}",
                          f"Otimista: {n(S['Otimista']['subs5'])} assinantes, {mi(S['Otimista']['gross5'])}",
                          "Gasto em marketing condicionado a marcos de conversão."], 16)
footer(s, k)

# 11. A rodada
k += 1
s = slide(TERRA)
text(s, 2, 1.4, 30, 1, "A RODADA", 13, RGBColor(0xFB, 0xED, 0xD6), True)
text(s, 2, 2.4, 30, 3, f"{mi(RAISE)} pré-seed · 18 meses", 44, WHITE, True)
uses = [("45%", "Marketing e aquisição"), ("35%", "Time"), ("10%", "IA, infra e apps"), ("10%", "Jurídico, LGPD e marca")]
for i, (p, t) in enumerate(uses):
    x = 2 + i * 7.6
    card(s, x, 7.2, 7, 4.6, RGBColor(0xC9, 0x5B, 0x44))
    text(s, x + 0.6, 7.7, 6, 1.8, p, 34, WHITE, True)
    text(s, x + 0.6, 9.7, 6, 2, t, 15, WHITE)
text(s, 2, 13, 30, 3, ["Marcos: lançamento público iOS e Android · conversão ≥ 30% · ~" + n(ys[1]['subs_end'] / 1000)
                       + " mil assinantes no ano 2 · primeiro estudo patrocinado"], 16, WHITE)

# 12. Time + contato
k += 1
s = slide()
title(s, "Time", "[Fundadores — preencher]")
text(s, 2, 6.5, 30, 6, ["[Nome · papel · experiência relevante]", "[Nome · papel · experiência relevante]",
                        "[Por que este time vence neste mercado]"], 18, MUTED)
text(s, 2, 14.5, 30, 1.5, "vico.bitrilha.com.br · [e-mail de contato]", 18, TERRA, True)
footer(s, k)

prs.save("pitch-vico.pptx")
print("pptx ok", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
