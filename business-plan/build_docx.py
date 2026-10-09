"""Gera plano-de-negocios-vico.docx (identidade Viço) a partir de model.py + data.json + charts/."""
import json
from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import model

TERRA, INK, MUTED, SOL = RGBColor(0xB8, 0x46, 0x2F), RGBColor(0x2A, 0x1B, 0x14), RGBColor(0x6E, 0x5C, 0x51), RGBColor(0xF2, 0xB6, 0x6D)
data = json.load(open("data.json", encoding="utf-8"))
ys = model.yearly(model.run())
ue = model.unit_economics()
S = data["scenarios"]
RAISE = 1_500_000


def brl(v, dec=0):
    s = f"{v:,.{dec}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {s}"


def mi(v):
    sign, a = ("-" if v < 0 else ""), abs(v)
    return f"{sign}R$ {a / 1e6:.1f} mi".replace(".", ",") if a >= 1e6 else f"{sign}R$ {a / 1e3:.0f} mil"


def n(v):
    return f"{v:,.0f}".replace(",", ".")


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(2)

base = doc.styles["Normal"]
base.font.name, base.font.size, base.font.color.rgb = "Calibri", Pt(11), INK
base.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
base.paragraph_format.space_after = Pt(6)
base.paragraph_format.line_spacing = 1.15
for lvl, size in ((1, 20), (2, 14), (3, 12)):
    st = doc.styles[f"Heading {lvl}"]
    st.font.name, st.font.size, st.font.bold, st.font.color.rgb = "Calibri", Pt(size), True, TERRA if lvl < 3 else INK
    st.element.rPr.rFonts.set(qn("w:asciiTheme"), "")
    st.paragraph_format.space_before = Pt(18 if lvl == 1 else 12)
    st.paragraph_format.space_after = Pt(6)


def shade(cell, hex_fill):
    tc = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"), sh.set(qn("w:color"), "auto"), sh.set(qn("w:fill"), hex_fill)
    tc.append(sh)


def table(headers, rows, widths=None, first_bold=False):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold, r.font.color.rgb, r.font.size = True, RGBColor(0xFF, 0xFF, 0xFF), Pt(10)
        shade(c, "B8462F")
    for k, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(10)
            r.bold = first_bold and i == 0
            if i > 0 and isinstance(v, str) and (v.startswith("R$") or v[:1].isdigit() or v.startswith("-")):
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if k % 2:
                shade(cells[i], "F6EEE7")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def para(text, bold=False, color=None, size=None, italic=False, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if color:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    if align:
        p.alignment = align
    return p


def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            r = p.add_run(it[0])
            r.bold = True
            p.add_run(" " + it[1])
        else:
            p.add_run(it)


def callout(text):
    t = doc.add_table(rows=1, cols=1)
    c = t.rows[0].cells[0]
    shade(c, "F9E6DE")
    c.text = ""
    r = c.paragraphs[0].add_run(text)
    r.font.color.rgb, r.bold = INK, True
    doc.add_paragraph()


def chart(name, width=16):
    doc.add_picture(f"charts/{name}.png", width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


# ---------- Capa ----------
for _ in range(6):
    doc.add_paragraph()
para("viço", bold=True, color=TERRA, size=72)
para("Plano de negócios", bold=True, size=28)
para("Sua pele lida. Seus produtos entendidos. Sua evolução medida.", color=MUTED, size=14)
for _ in range(10):
    doc.add_paragraph()
para(f"Versão para investidores · {date.today().strftime('%m/%Y')}", color=MUTED)
para("Confidencial — não distribuir sem autorização.", color=MUTED, italic=True, size=9)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---------- 1. Sumário executivo ----------
doc.add_heading("1. Sumário executivo", 1)
para("O Viço é um app que lê a pele por selfie, lê os produtos pelo rótulo e monta a rotina na ordem certa — "
     "e mede, semana a semana, se aquilo está funcionando. É o único produto no Brasil que fecha esse ciclo: "
     "hoje o consumidor precisa de um app para analisar a pele, outro para entender rótulos e uma planilha para acompanhar resultado.")
callout(f"Pedimos {mi(RAISE)} em rodada pré-seed para 18 meses de operação: lançar ao público, validar o funil pago "
        f"e chegar a ~{n(ys[1]['subs_end'] / 1000)} mil assinantes no fim do ano 2 (cenário base).")
table(["Indicador (cenário base)", "Ano 1", "Ano 3", "Ano 5"], [
    ["Assinantes ativos (fim do ano)", n(ys[0]["subs_end"]), n(ys[2]["subs_end"]), n(ys[4]["subs_end"])],
    ["Receita bruta", mi(ys[0]["gross"]), mi(ys[2]["gross"]), mi(ys[4]["gross"])],
    ["EBITDA", mi(ys[0]["ebitda"]), mi(ys[2]["ebitda"]), mi(ys[4]["ebitda"])],
], widths=[7, 3, 3, 3], first_bold=True)
bullets([
    ("Produto no ar:", "beta fechado rodando em vico.bitrilha.com.br, app Android gerado, iOS em seguida. Análise por IA, leitura de rótulo, motor de rotina com detector de conflitos, evolução e pagamentos prontos."),
    ("Mercado:", "Brasil é o 3º maior mercado de beleza do mundo (~R$ 200 bi em 2025, ABIHPEC); dermocosméticos somaram R$ 12,3 bi em 12 meses, +8,3% (IQVIA)."),
    ("Prova de categoria:", "o Cal AI, que aplicou a mesma mecânica (foto → IA → acompanhamento) a calorias, chegou a ~US$ 30 mi de ARR sem capital externo e foi comprado pela MyFitnessPal (imprensa)."),
    ("Modelo:", f"assinatura B2C (R$ 24,90/mês ou R$ 199/ano, 7 dias grátis) + receitas B2B a partir do ano 2. LTV de contribuição ~{brl(ue['ltv'])} por assinante."),
])

# ---------- 2. Problema ----------
doc.add_heading("2. O problema", 1)
para("Analisamos 216 avaliações públicas de 9 apps de skincare (Skan, Yuka, FeelinMySkin, Mary Kay Skin Analyzer, SkinSAFE e outros). "
     "Os incômodos se repetem:")
table(["Dor", "Avaliações", "Apps"], [
    ["Câmera/foto falha no scan", "20", "8"],
    ["Paywall sem valor percebido", "16", "4"],
    ["Querem acompanhar a evolução ligada aos produtos", "16", "3"],
    ["Rotina rígida (frequência, passo esquecido)", "13", "4"],
    ["Produto não encontrado / base fraca fora dos EUA", "10", "5"],
    ["Resultado inconsistente entre fotos", "9", "4"],
    ["Cobrança e cancelamento difíceis", "8", "3"],
], widths=[10, 3, 3])
para("Por trás disso há um problema maior: quem compra skincare não sabe se está funcionando. Os apps existentes são ou "
     "scanner de rosto, ou scanner de rótulo, ou diário — nenhum liga os três. E as ferramentas das grandes marcas "
     "(como a de análise de pele da Natura, de 2022) recomendam apenas o portfólio da própria marca.")

# ---------- 3. Solução ----------
doc.add_heading("3. A solução: Viço", 1)
bullets([
    ("Selfie → Skin Score.", "Nota de 0 a 100 e 7 métricas (textura, vermelhidão, poros, linhas, manchas, oleosidade, hidratação), por zona do rosto. A câmera só dispara com luz e enquadramento bons, o que torna as fotos comparáveis."),
    ("Rótulo → ingredientes.", "Foto do verso do frasco; a IA lê a lista INCI de qualquer marca, inclusive de farmácia brasileira."),
    ("Rotina na ordem certa.", "Motor determinístico (sem IA, auditável) ordena manhã e noite, distribui ativos na semana e avisa conflitos — com correção em um toque."),
    ("Evolução.", "Gráfico por métrica, antes/depois e card compartilhável sem rosto (crescimento orgânico em redes)."),
    ("Confiança.", "Fotos em armazenamento privado, nunca usadas para treinar IA, consentimento LGPD explícito; linguagem cosmética, sem diagnóstico; recomendação neutra, patrocínio sempre marcado."),
])
para("Status: produto construído (Next.js, Supabase, Stripe, IA de visão), beta fechado no ar, app Android empacotado, painel de acessos para testers.", italic=True, color=MUTED)

# ---------- 4. Mercado ----------
doc.add_heading("4. Mercado", 1)
table(["Recorte", "Tamanho", "Fonte"], [
    ["Beleza e cuidados pessoais — Brasil", "~R$ 200 bi (2025); 3º maior do mundo", "ABIHPEC"],
    ["Beleza — Brasil (varejo)", "R$ 186,9 bi (2025), +6,8%", "Euromonitor"],
    ["Dermocosméticos — Brasil", "R$ 12,3 bi (12m até jul/25), +8,3%", "IQVIA"],
    ["Apps de análise de pele — global", "US$ 1,82 bi (2025), CAGR 17,2%", "Future Market Insights"],
], widths=[6, 6, 4])
bullets([
    "88% das mulheres e 64% dos homens usam produtos de skincare (MindMiners, 2024).",
    "23% dos brasileiros já usam assistentes de IA para pesquisar ou escolher beleza; influenciadores respondem por 32% da descoberta de produtos (Mercado & Consumo, out/2026).",
    "83% das vendas de dermocosméticos ainda acontecem em loja física — o consumidor pesquisa no celular e compra na farmácia: espaço para afiliados e white-label com redes de farmácia.",
])
callout(f"Meta do ano 5 (cenário base): {mi(ys[4]['gross'])} de receita — cerca de "
        f"{ys[4]['gross'] / 12.3e9 * 100:.2f}% do mercado anual de dermocosméticos no Brasil.".replace(".", ",", 1))

# ---------- 5. Concorrência ----------
doc.add_heading("5. Concorrência", 1)
table(["", "Lê a pele", "Lê o rótulo", "Monta rotina", "Mede evolução", "Neutro"], [
    ["Viço", "Sim", "Sim", "Sim, com conflitos", "Sim", "Sim"],
    ["Skan", "Sim", "—", "Parcial", "Parcial", "—"],
    ["Yuka", "—", "Sim", "—", "—", "Sim"],
    ["FeelinMySkin", "—", "Parcial", "Diário", "Sim", "Sim"],
    ["Ferramentas de marca (Natura, Mary Kay)", "Sim", "—", "Só a marca", "—", "—"],
], widths=[4.5, 2, 2, 2.6, 2.4, 1.8], first_bold=True)
para("Vantagens defensáveis: base de ingredientes e regras de conflito em português, histórico longitudinal do usuário "
     "(quanto mais tempo no app, mais valor), e dados agregados e anonimizados de eficácia real de ativos — inexistentes no mercado brasileiro.")

# ---------- 6. Modelo de negócio ----------
doc.add_heading("6. Modelo de negócio", 1)
table(["Fonte de receita", "Como", "Quando"], [
    ["Assinatura B2C", "R$ 24,90/mês ou R$ 199/ano (R$ 16,58/mês), 7 dias grátis", "Lançamento"],
    ["Afiliados", "Link de recompra quando o produto acaba (comissão por venda)", "Ano 1"],
    ["Estudos de eficácia patrocinados", "Marcas pagam para medir produto com usuários voluntários (métricas anônimas)", "Ano 2"],
    ["Painel de tendências", "Dados agregados e anonimizados de pele e uso de ativos, por região", "Ano 3"],
    ["White-label para farmácias", "Análise e rotina dentro do app da rede", "Ano 3+"],
], widths=[4.5, 8.5, 3])
doc.add_heading("Unit economics (cenário base)", 3)
table(["Métrica", "Valor"], [
    ["ARPU bruto mensal", brl(ue["arpu"], 2)],
    ["ARPU líquido (taxas, impostos e IA)", brl(ue["net_arpu"], 2)],
    ["Vida média do assinante", f"{ue['life_months']:.1f} meses".replace(".", ",")],
    ["LTV de contribuição", brl(ue["ltv"], 2)],
    ["CAC pago — ano 1", brl(ys[0]["cac"], 2)],
    ["CAC misto — ano 1 (com orgânico)", brl(ys[0]["cac"] / (1 + model.Y["organic_ratio"][0]), 2)],
    ["Custo de IA por assinante/mês", brl(model.P["ai_cost_per_sub"], 2)],
], widths=[9, 5])

# ---------- 7. Go-to-market ----------
doc.add_heading("7. Go-to-market", 1)
bullets([
    ("Fase 1 — beta fechado (agora):", "testers convidados, calibrar funil, custo de IA e retenção."),
    ("Fase 2 — lançamento B2C:", "TikTok e Reels com o card \"minha pele em 30 dias\"; micro-influenciadores de skincare (5–20 mil seguidores) por permuta e afiliado; SEO de ingredientes (\"retinol e ácido pode junto?\") gerado da base de regras."),
    ("Fase 3 — B2B:", "estudos patrocinados com marcas de dermocosméticos; parcerias com redes de farmácia (afiliado e white-label)."),
])
chart("funil", 15)

# ---------- 8. Projeções ----------
doc.add_heading("8. Projeções financeiras", 1)
para("Modelo mensal de 60 meses, com fórmulas abertas na planilha anexa (modelo-financeiro-vico.xlsx). Premissas-chave: "
     "conversão de teste para pagante de 40% (mediana de apps de Saúde & Fitness, RevenueCat 2025), churn mensal de 12% no plano mensal, "
     "renovação anual de 35%, 60% dos novos assinantes no plano anual.", color=MUTED)
table(["", "Ano 1", "Ano 2", "Ano 3", "Ano 4", "Ano 5"], [
    ["Assinantes (fim)"] + [n(y["subs_end"]) for y in ys],
    ["Receita bruta"] + [mi(y["gross"]) for y in ys],
    ["  dos quais B2B"] + [mi(y["b2b"]) for y in ys],
    ["Receita líquida"] + [mi(y["net"]) for y in ys],
    ["Marketing"] + [mi(y["marketing"]) for y in ys],
    ["EBITDA"] + [mi(y["ebitda"]) for y in ys],
], widths=[4, 2.4, 2.4, 2.4, 2.4, 2.4], first_bold=True)
chart("receita")
chart("caixa")
doc.add_heading("Cenários", 3)
table(["", "Conservador", "Base", "Otimista"], [
    ["Assinantes no ano 5"] + [n(S[k]["subs5"]) for k in S],
    ["Receita no ano 5"] + [mi(S[k]["gross5"]) for k in S],
    ["EBITDA no ano 5"] + [mi(S[k]["ebitda5"]) for k in S],
    ["Necessidade de caixa"] + [mi(S[k]["need"]) for k in S],
], widths=[5, 3.5, 3.5, 3.5], first_bold=True)
para("Conservador: teste→pago 25%, churn 16%, renovação 25%. Otimista: 55%, 9%, 45%.", italic=True, color=MUTED, size=9)
chart("cenarios", 15)

# ---------- 9. Captação ----------
doc.add_heading("9. A rodada", 1)
callout(f"Pré-seed de {mi(RAISE)}, via mútuo conversível ou SAFE. Valuation a negociar.")
para(f"O valor cobre com folga a necessidade de caixa do cenário base ({mi(S['Base']['need'])}) e dá 18 meses de pista "
     f"para provar o funil antes de uma seed. No cenário conservador a necessidade sobe para {mi(S['Conservador']['need'])}: "
     "por isso o plano tem marcos de corte de gasto em marketing se a conversão ficar abaixo de 25%.")
table(["Uso dos recursos", "%", "Valor"], [
    ["Marketing e aquisição (testes de canal, influenciadores)", "45%", mi(RAISE * 0.45)],
    ["Time (produto, crescimento, dados)", "35%", mi(RAISE * 0.35)],
    ["IA, infraestrutura e app iOS/Android", "10%", mi(RAISE * 0.10)],
    ["Jurídico, LGPD, marca (INPI) e regulatório", "10%", mi(RAISE * 0.10)],
], widths=[10, 2, 3])
doc.add_heading("Marcos até a próxima rodada", 3)
bullets(["Lançamento público em iOS e Android.", "Conversão teste→pago ≥ 30% e renovação anual ≥ 30% medidas.",
         f"~{n(ys[1]['subs_end'] / 1000)} mil assinantes ativos no fim do ano 2 (base).", "Primeiro estudo patrocinado vendido."])

# ---------- 10. Riscos ----------
doc.add_heading("10. Riscos e mitigação", 1)
table(["Risco", "Mitigação"], [
    ["LGPD — foto do rosto é dado sensível", "Consentimento explícito e separado; armazenamento privado; exclusão em um toque; nunca treina IA; revisão jurídica antes do lançamento."],
    ["Regulatório (ANVISA) — parecer diagnóstico", "Linguagem estritamente cosmética; aviso de \"olhar profissional\" quando a leitura foge do normal; sem recomendação de medicamento."],
    ["Custo de IA por análise", "Limite diário por usuário; custo monitorado no gateway; troca de modelo sem mudar o app."],
    ["Grandes marcas lançarem ferramentas próprias", "Neutralidade é o diferencial: marca não recomenda concorrente. Dados longitudinais criam barreira."],
    ["Dependência das lojas de apps", "Web/PWA funcional e cobrança via web (Stripe) como canal paralelo."],
    ["Conversão e retenção abaixo do previsto", "Cenário conservador modelado; gasto em marketing condicionado a marcos."],
], widths=[5.5, 10.5])

# ---------- 11. Time ----------
doc.add_heading("11. Time", 1)
para("[Preencher: fundadores, papéis, experiência relevante e participação. Investidores pré-seed decidem muito pelo time — "
     "incluir histórico de produtos lançados e por que este time vence neste mercado.]", italic=True, color=MUTED)

# ---------- Fontes ----------
doc.add_heading("Fontes", 1)
for dado, fonte, link in data["sources"]:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(f"{dado} — {fonte}. ").font.size = Pt(9)
    r = p.add_run(link)
    r.font.size, r.font.color.rgb = Pt(9), MUTED
para("Projeções são estimativas baseadas em premissas declaradas e benchmarks públicos; não constituem garantia de resultado.",
     italic=True, color=MUTED, size=9)

# rodapé
footer = sec.footer.paragraphs[0]
footer.text = "viço · plano de negócios · confidencial"
footer.runs[0].font.size, footer.runs[0].font.color.rgb = Pt(8), MUTED

doc.save("plano-de-negocios-vico.docx")
print("docx ok")
